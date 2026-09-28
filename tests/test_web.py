import io
import shutil
import time
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from edugen import ingest
from edugen.web import create_app

ROOT = Path(__file__).resolve().parents[1]
SAMPLE_PDF = ROOT / "products" / "sample-ko5-1" / "source.pdf"


def _pptx_bytes() -> bytes:
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    prs = pptx.Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for t in ("3단원 분수의 덧셈", "활동 1. 피자 나누기", "활동 2. 수직선", "정답"):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = t
        s.placeholders[1].text = "초등 4학년 수학 (漢字)"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


@pytest.fixture()
def client(tmp_path):
    return TestClient(create_app(tmp_path)), tmp_path


def _wait(c, job_id, timeout=240):
    end = time.time() + timeout
    while time.time() < end:
        j = c.get(f"/api/jobs/{job_id}").json()
        if j["state"] in ("done", "failed"):
            return j
        time.sleep(1)
    raise AssertionError("job timeout")


def test_index(client):
    c, _ = client
    r = c.get("/")
    assert r.status_code == 200 and "자료 파일 올리기" in r.text


@pytest.mark.parametrize("name,data,msg", [
    ("a.docx", b"PK\x03\x04xxxx", "PDF, PPTX, PPT"),
    ("fake.pdf", b"hello world", "PDF 또는 PPT/PPTX"),
    ("fake.pptx", b"PK\x03\x04broken", "PPTX 파일이 아니거나"),
    ("empty.pdf", b"", "빈 파일"),
])
def test_rejects_bad_files_and_cleans_up(client, name, data, msg):
    c, root = client
    r = c.post("/api/upload", files={"file": (name, data)})
    assert r.status_code == 400
    assert msg in r.json()["error"]
    assert [p for p in root.iterdir() if p.is_dir()] == []


def test_rejects_bad_sku(client):
    c, _ = client
    r = c.post("/api/upload", files={"file": ("a.pdf", SAMPLE_PDF.read_bytes())}, data={"sku": "../x"})
    assert r.status_code == 400


def test_detect_kind(tmp_path):
    p = tmp_path / "x.bin"
    p.write_bytes(_pptx_bytes())
    assert ingest.detect_kind(p) == ".pptx"
    assert ingest.detect_kind(SAMPLE_PDF) == ".pdf"


def test_file_route_blocks_traversal(client):
    c, root = client
    (root / "abc" / "out").mkdir(parents=True)
    (root / "abc" / "meta.yaml").write_text("sku: abc")
    assert c.get("/p/abc/files/../meta.yaml").status_code == 404
    assert c.get("/p/abc/files/%2E%2E/meta.yaml").status_code == 404


needs_render = pytest.mark.skipif(ingest.find_soffice() is None, reason="LibreOffice 필요")


@needs_render
@pytest.mark.parametrize("filename,kind", [("활동지.pdf", "pdf"), ("수학 교안.pptx", "pptx")])
def test_upload_end_to_end(client, filename, kind):
    c, root = client
    data = SAMPLE_PDF.read_bytes() if kind == "pdf" else _pptx_bytes()
    r = c.post("/api/upload", files={"file": (filename, data)},
               data={"store_name": "테스트", "subject": "수학", "sku": f"t-{kind}", "has_answer_key": "yes"})
    assert r.status_code == 200, r.text
    j = _wait(c, r.json()["job_id"])
    assert j["state"] == "done", j
    out = root / f"t-{kind}" / "out"
    assert (out / "thumb_main.jpg").exists() and list(out.glob("detail_*.jpg"))
    assert (root / f"t-{kind}" / f"source.{kind}").exists()
    page = c.get(f"/p/t-{kind}")
    assert page.status_code == 200 and "상세페이지" in page.text
    z = c.get(f"/p/t-{kind}/download.zip")
    assert z.status_code == 200 and z.content[:2] == b"PK"
    if kind == "pptx":
        import pymupdf
        fonts = [f[3] for f in pymupdf.open(root / "t-pptx" / ingest.CONVERTED_NAME)[0].get_fonts()]
        assert not any("WenQuanYi" in f for f in fonts), fonts
        assert "가로형" in "\n".join(j["logs"])
    # 같은 코드로 다른 형식을 다시 올리면 이전 원본을 지운다
    other = ("b.pptx", _pptx_bytes()) if kind == "pdf" else ("b.pdf", SAMPLE_PDF.read_bytes())
    r2 = c.post("/api/upload", files={"file": other}, data={"sku": f"t-{kind}"})
    assert _wait(c, r2.json()["job_id"])["state"] == "done"
    sources = sorted(p.name for p in (root / f"t-{kind}").glob("source.*") if p.name != ingest.CONVERTED_NAME)
    assert len(sources) == 1 and sources[0] != f"source.{kind}"
